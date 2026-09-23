from pathlib import Path
import csv
import re
import unicodedata

import pdfplumber
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from langchain_core.documents import Document


# --------------------------------------------------
# Supported file types
# --------------------------------------------------

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".txt",
    ".md",
    ".docx",
    ".html",
    ".htm",
    ".csv",
}


# --------------------------------------------------
# Text Cleaning
# --------------------------------------------------

def clean_text(text):
    """
    Normalize and clean extracted text.
    """

    if not text:
        return ""

    text = unicodedata.normalize(
        "NFKC",
        text
    )

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


# --------------------------------------------------
# Table Formatting
# --------------------------------------------------

def format_table(table):
    """
    Convert a table into Markdown format.
    """

    if not table:
        return ""

    rows = []

    for row in table:

        cleaned_row = [
            clean_text(cell or "")
            for cell in row
        ]

        rows.append(cleaned_row)

    if not rows:
        return ""

    header = rows[0]

    if not header:
        return ""

    markdown = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(
            ["---"] * len(header)
        ) + " |"
    ]

    for row in rows[1:]:

        # Make row length equal to header
        row = row + [""] * (
            len(header) - len(row)
        )

        row = row[:len(header)]

        markdown.append(
            "| " + " | ".join(row) + " |"
        )

    return "\n".join(markdown)


# --------------------------------------------------
# PDF Loader
# --------------------------------------------------

def load_pdf(file_path):
    """
    Extract text and tables from PDF.

    PDF pages are stored internally using
    zero-based page numbers.
    """

    documents = []

    with pdfplumber.open(file_path) as pdf:

        for page_number, page in enumerate(
            pdf.pages
        ):

            parts = []

            # Extract normal text
            text = page.extract_text()

            if text:
                text = clean_text(text)

                if text:
                    parts.append(text)

            # Extract tables
            tables = page.extract_tables()

            for table in tables:

                table_text = format_table(
                    table
                )

                if table_text:
                    parts.append(table_text)

            content = "\n\n".join(
                parts
            ).strip()

            if not content:
                continue

            documents.append(
                Document(
                    page_content=content,
                    metadata={
                        "source": str(file_path),
                        "page": page_number,
                        "file_type": "pdf",
                    },
                )
            )

    return documents


# --------------------------------------------------
# DOCX Loader
# --------------------------------------------------

def load_docx(file_path):
    """
    Extract paragraphs and tables from DOCX.
    """

    # python-docx Document
    doc = DocxDocument(file_path)

    parts = []

    # ------------------------------
    # Paragraphs
    # ------------------------------

    for paragraph in doc.paragraphs:

        text = clean_text(
            paragraph.text
        )

        if text:
            parts.append(text)

    # ------------------------------
    # Tables
    # ------------------------------

    for table in doc.tables:

        rows = []

        for row in table.rows:

            rows.append(
                [
                    clean_text(
                        cell.text
                    )
                    for cell in row.cells
                ]
            )

        table_text = format_table(
            rows
        )

        if table_text:
            parts.append(table_text)

    content = "\n\n".join(
        parts
    ).strip()

    if not content:
        return []

    return [
        Document(
            page_content=content,
            metadata={
                "source": str(file_path),
                "file_type": "docx",
            },
        )
    ]


# --------------------------------------------------
# HTML Loader
# --------------------------------------------------

def load_html(file_path):
    """
    Extract text and tables from HTML.
    """

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore",
    ) as file:

        html = file.read()

    soup = BeautifulSoup(
        html,
        "html.parser"
    )

    # Remove scripts and styles
    for element in soup(
        ["script", "style"]
    ):
        element.decompose()

    parts = []

    # ------------------------------
    # Normal text
    # ------------------------------

    text = soup.get_text(
        separator=" ",
        strip=True
    )

    text = clean_text(text)

    if text:
        parts.append(text)

    # ------------------------------
    # Tables
    # ------------------------------

    for table in soup.find_all("table"):

        rows = []

        for tr in table.find_all("tr"):

            cells = tr.find_all(
                ["th", "td"]
            )

            row = [
                clean_text(
                    cell.get_text(
                        separator=" ",
                        strip=True
                    )
                )
                for cell in cells
            ]

            if row:
                rows.append(row)

        table_text = format_table(
            rows
        )

        if table_text:
            parts.append(table_text)

    content = "\n\n".join(
        parts
    ).strip()

    if not content:
        return []

    return [
        Document(
            page_content=content,
            metadata={
                "source": str(file_path),
                "file_type": "html",
            },
        )
    ]


# --------------------------------------------------
# TXT / Markdown Loader
# --------------------------------------------------

def load_text(file_path):
    """
    Load TXT or Markdown files.
    """

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore",
    ) as file:

        content = file.read()

    content = clean_text(content)

    if not content:
        return []

    file_type = (
        file_path
        .suffix
        .lower()
        .lstrip(".")
    )

    return [
        Document(
            page_content=content,
            metadata={
                "source": str(file_path),
                "file_type": file_type,
            },
        )
    ]


# --------------------------------------------------
# CSV Loader
# --------------------------------------------------

def load_csv(file_path):
    """
    Load CSV and convert it into Markdown table.
    """

    rows = []

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore",
        newline="",
    ) as file:

        reader = csv.reader(file)

        for row in reader:

            rows.append(
                [
                    clean_text(cell)
                    for cell in row
                ]
            )

    if not rows:
        return []

    content = format_table(
        rows
    )

    if not content:
        return []

    return [
        Document(
            page_content=content,
            metadata={
                "source": str(file_path),
                "file_type": "csv",
            },
        )
    ]


# --------------------------------------------------
# Single Document Loader
# --------------------------------------------------

def load_document(file_path):
    """
    Load one supported document.
    """

    file_path = Path(file_path)

    extension = file_path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:

        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    if extension == ".pdf":

        return load_pdf(file_path)

    elif extension == ".docx":

        return load_docx(file_path)

    elif extension in {".html", ".htm"}:

        return load_html(file_path)

    elif extension in {".txt", ".md"}:

        return load_text(file_path)

    elif extension == ".csv":

        return load_csv(file_path)

    return []
